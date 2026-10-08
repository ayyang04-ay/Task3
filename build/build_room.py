"""Assemble the WU397 data room from the eleven source-data workbooks.

Filenames are assay-keyed, never figure-keyed: the source sheets ship names
like 'Fig 5d and e', which would telegraph which paper figure a table became.
"""
import sys, csv, os, re
import openpyxl

SRC, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)
NBSP = "\xa0"

def clean(v):
    return v.replace(NBSP, " ").strip() if isinstance(v, str) else v

def load(n, sheet):
    wb = openpyxl.load_workbook(f"{SRC}/41590_2026_2569_MOESM{n}_ESM.xlsx", data_only=True)
    return [[clean(c) for c in row] for row in wb[sheet].iter_rows(values_only=True)]

def arm_cols(rows, hdr_row=0, id_row=1):
    """row0 carries sparse arm-group labels, row1 the participant ids."""
    arms, cur = {}, None
    for j, v in enumerate(rows[hdr_row]):
        if v not in (None, ""): cur = str(v)
        arms[j] = cur
    ids = {j: (str(v) if v not in (None, "") else None)
           for j, v in enumerate(rows[id_row])}
    return arms, ids

def write(name, header, records):
    p = f"{OUT}/{name}"
    with open(p, "w", newline="") as fh:
        w = csv.writer(fh); w.writerow(header); w.writerows(records)
    print(f"  {name:38s} {len(records):8,d} rows")

# ---------------------------------------------------------------- serology
# MOESM4 'Fig 1d': four stacked EC50 blocks then four fold-change blocks.
rows = load(4, "Fig 1d")
arms, ids = arm_cols(rows)
BLOCKS = {"A/H1": 0, "A/H3": 7, "B/Yam": 14, "B/Vic": 21}
FCB    = {"A/H1": 28, "A/H3": 33, "B/Yam": 38, "B/Vic": 43}
rec = []
for ag, r0 in BLOCKS.items():
    for ri in range(r0 + 2, r0 + 6):
        if ri >= len(rows) or rows[ri][0] in (None, ""): continue
        wk = rows[ri][0]
        for j in range(1, len(rows[ri])):
            if ids.get(j) and isinstance(rows[ri][j], (int, float)):
                rec.append([ag, f"wk{wk}", arms[j], ids[j], "ec50", rows[ri][j]])
for ag, r0 in FCB.items():
    for ri in range(r0 + 2, r0 + 4):
        if ri >= len(rows) or rows[ri][0] in (None, ""): continue
        wk = rows[ri][0]
        for j in range(1, len(rows[ri])):
            if ids.get(j) and isinstance(rows[ri][j], (int, float)):
                rec.append([ag, f"wk{wk}", arms[j], ids[j], "fold_change", rows[ri][j]])
write("elisa_igg_titers.csv",
      ["antigen", "timepoint", "arm", "participant", "measure", "value"], rec)

# MOESM9 'XData Fig1c': HAI, four stacked antigen blocks.
rows = load(9, "XData Fig1c")
arms, ids = arm_cols(rows)
rec = []
for ag, r0 in {"A/H1N1": 0, "A/H3N2": 7, "B/Yam": 14, "B/Vic": 21}.items():
    for ri in range(r0 + 2, r0 + 6):
        if ri >= len(rows) or rows[ri][0] in (None, ""): continue
        wk = rows[ri][0]
        for j in range(1, len(rows[ri])):
            if ids.get(j) and isinstance(rows[ri][j], (int, float)):
                rec.append([ag, f"wk{wk}", arms[j], ids[j], rows[ri][j]])
write("hai_titers.csv",
      ["antigen", "timepoint", "arm", "participant", "hai_titer"], rec)

# MOESM9 1a/1b: IgM, distractor.
rows = load(9, "XData Fig1b")
arms, ids = arm_cols(rows)
rec = [[rows[ri][0], arms[j], ids[j], rows[ri][j]]
       for ri in range(2, len(rows)) if rows[ri][0] not in (None, "")
       for j in range(1, len(rows[ri]))
       if ids.get(j) and isinstance(rows[ri][j], (int, float))]
write("igm_fold_change.csv", ["antigen", "arm", "participant", "fold_change"], rec)

# ------------------------------------------------- functional: neutralization
wb = openpyxl.load_workbook(f"{SRC}/41590_2026_2569_MOESM8_ESM.xlsx", data_only=True)
ws = wb["Fig 5d and e"]
rec = [[clean(c) for c in row] for row in ws.iter_rows(values_only=True)
       if row and any(v not in (None, "") for v in row)]
hdr, body = rec[0], rec[1:]
write("neutralization_seqbased.csv", hdr, body)

# binding breadth panels
for sheet, name in (("Fig 5b", "binding_breadth_h1.csv"),
                    ("Fig 5c", "binding_breadth_h3.csv")):
    rows = load(8, sheet)
    arms, ids = arm_cols(rows)
    rec = [[rows[ri][0], arms[j], ids[j], rows[ri][j]]
           for ri in range(2, len(rows)) if rows[ri][0] not in (None, "")
           for j in range(1, len(rows[ri]))
           if ids.get(j) and isinstance(rows[ri][j], (int, float))]
    write(name, ["strain", "arm", "participant", "log2_fold_change_mfi"], rec)

# conventional microneutralization
rec = []
for sheet, kind, meas in (("XData Fig 7a", "homologous", "titer"),
                          ("XData Fig 7b", "heterologous", "titer"),
                          ("XData Fig 7c", "homologous", "fold_change"),
                          ("XData Fig 7d", "heterologous", "fold_change")):
    rows = load(14, sheet)
    arms, ids = arm_cols(rows)
    for ri in range(2, len(rows)):
        if rows[ri][0] in (None, ""): continue
        wk = rows[ri][0]
        for j in range(1, len(rows[ri])):
            if ids.get(j) and isinstance(rows[ri][j], (int, float)):
                rec.append([kind, f"wk{wk}", arms[j], ids[j], meas, rows[ri][j]])
write("microneutralization.csv",
      ["panel", "timepoint", "arm", "participant", "measure", "value"], rec)

# monoclonal antibody binding panels
rec = []
for sheet, grp in (("XData Fig 6c", "group_I"), ("XData Fig 6d", "group_II")):
    rows = load(13, sheet)
    arms, ids = arm_cols(rows)
    for ri in range(2, len(rows)):
        if rows[ri][0] in (None, ""): continue
        for j in range(1, len(rows[ri])):
            if ids.get(j) and isinstance(rows[ri][j], (int, float)):
                rec.append([grp, rows[ri][0], arms[j], ids[j], rows[ri][j]])
write("mab_binding_panels.csv",
      ["panel", "antigen", "arm", "mab_id", "fold_over_control"], rec)

# ----------------------------------------------------- cellular / mechanistic
# MOESM6 'Fig 3d' carries arm labels; MOESM11 'Fig 4b' does not and is dropped.
rows = load(6, "Fig 3d")
arms, ids = arm_cols(rows)
rec = [[f"wk{rows[ri][0]}", arms[j], ids[j], rows[ri][j]]
       for ri in range(2, len(rows)) if rows[ri][0] not in (None, "")
       for j in range(1, len(rows[ri]))
       if ids.get(j) and isinstance(rows[ri][j], (int, float))]
write("gc_frequencies.csv",
      ["timepoint", "arm", "participant", "pct_ha_specific_igg_gc"], rec)

rows = load(5, "Fig 2b")
arms, ids = arm_cols(rows)
rec = [[f"wk{rows[ri][0]}", arms[j], ids[j], rows[ri][j]]
       for ri in range(2, len(rows)) if rows[ri][0] not in (None, "")
       for j in range(1, len(rows[ri]))
       if ids.get(j) and isinstance(rows[ri][j], (int, float))]
write("plasmablast_frequencies.csv",
      ["timepoint", "arm", "participant", "pct_plasmablast"], rec)

rows = load(6, "Fig 3a")
write("pb_gc_clonal_overlap.csv", [str(c) for c in rows[0]],
      [r for r in rows[1:] if r[0] not in (None, "")])

rows = load(7, "Fig 4c")
arms, ids = arm_cols(rows, 1, 1)
rec = []
for ri in range(2, len(rows)):
    if rows[ri][0] in (None, ""): continue
    for j in range(1, len(rows[ri])):
        if isinstance(rows[ri][j], (int, float)):
            rec.append([f"wk{rows[ri][0]}", arms[j], rows[ri][j]])
write("serum_igg_clonotypes.csv", ["timepoint", "arm", "n_clonotypes"], rec)

rec = []
for sheet, metric in (("XData Fig 5a", "pct_pre_existing_abundance"),
                      ("XData Fig 5b", "n_pre_existing_clonotypes"),
                      ("XData Fig 5c", "n_vaccine_elicited_clonotypes")):
    rows = load(12, sheet)
    for ri in range(2, len(rows)):
        for j, arm in ((0, "mRNA-1010"), (1, "Fluarix")):
            if j < len(rows[ri]) and isinstance(rows[ri][j], (int, float)):
                rec.append([metric, arm, rows[ri][j]])
write("repertoire_clonotype_counts.csv", ["metric", "arm", "value"], rec)

rows = load(12, "XData Fig 5f")
arms, ids = arm_cols(rows, 1, 1)
rec = []
for ri in range(2, len(rows)):
    if rows[ri][0] in (None, ""): continue
    for j in range(1, len(rows[ri])):
        if isinstance(rows[ri][j], (int, float)):
            rec.append([f"wk{rows[ri][0]}", arms[j], rows[ri][j]])
write("cdrh3_diversity_index.csv", ["timepoint", "arm", "diversity_index"], rec)
